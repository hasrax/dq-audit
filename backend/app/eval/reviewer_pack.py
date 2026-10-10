"""Second-reviewer packs (Section 5): triage agreement + recall check sheets.

Usage (from backend/):
    python -m app.eval.reviewer_pack --reviewer R2 [--seed 20261014]
        [--companies "Company A,Company B"] [--data-root ../data]
        [--out ../evaluation/reviewer_packs]

Writes Triage_Agreement_<R>.xlsx and Recall_Check_<R>.xlsx. The content
is identical for every reviewer (same seed); only the file name and the
label differ. The first run also writes Recall_Check_MHG.xlsx (all 200
rows, with the reviewer subset marked). Answer keys go to
evaluation/reviewer_keys/ and are never shown to reviewers; a later run
whose keys would differ (e.g. the data changed) is refused.

Triage sheet: 15 rule + 15 AI issues (all statuses) from each company's
latest audit, shuffled. Source, decisions, notes and reviewer names are
not shown.

Recall sheet: rows that no issue points at -- a systematic random sample,
stratified by table row share, minimum 10 per table where possible.
"Unflagged" follows the majority rule agreed on 10 Oct 2026: an issue
that covers more than half of a table's rows is a column-wide finding; it
is listed on the table's tab as already reported and does not mark rows.

Values are masked exactly as in the hosted copy (sync_cloud.Masker).
Google Sheets compatible: no formulas or macros, list dropdowns only.
"""
from __future__ import annotations

import argparse
import csv
import random
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from psycopg.rows import dict_row

from app.aliases import client_label
from app.db import repository
from app.db.sync_cloud import Masker
from app.rules.engine import run_rules
from app.schema_docs.store import display_aliases
from app import workspaces as ws

ROOT = Path(__file__).resolve().parents[3]
OUT_DEFAULT = ROOT / "evaluation" / "reviewer_packs"
KEYS_DIR = ROOT / "evaluation" / "reviewer_keys"
N_TRIAGE, N_PER_SOURCE = 30, 15
N_RECALL, N_SUBSET, MIN_PER_TABLE = 200, 50, 10
COLUMN_WIDE = 0.5            # an issue covering more than this share of a table is column-wide
N_EXAMPLES = 3
DEADLINE = "Tuesday 13 October 2026"
DECISIONS = ["Accept", "Reject", "Unsure"]
DIMENSIONS = ["Completeness", "Validity", "Consistency", "Uniqueness", "Timeliness", "Cross-table"]
FILL = PatternFill("solid", start_color="FFF6E7")          # cells the reviewer fills in
HEAD_FILL = PatternFill("solid", start_color="E8EEF4")
WRAP = Alignment(wrap_text=True, vertical="top")
BOLD = Font(bold=True)


# --- sampling helpers (pure, tested) -------------------------------------------

def allocate(sizes: dict[str, int], total: int, minimum: int = 0) -> dict[str, int]:
    """Split `total` across groups in proportion to `sizes` (largest remainder),
    giving each group at least `minimum` where it has that many items."""
    sizes = {k: v for k, v in sizes.items() if v > 0}
    total = min(total, sum(sizes.values()))
    alloc = {k: min(minimum, v) for k, v in sizes.items()}
    left = total - sum(alloc.values())
    if left < 0:                        # minimums alone exceed the total: proportional only
        alloc, left = {k: 0 for k in sizes}, total
    room = {k: sizes[k] - alloc[k] for k in sizes}
    while left > 0 and any(room.values()):
        weight = sum(sizes[k] for k in sizes if room[k])
        shares = {k: left * sizes[k] / weight for k in sizes if room[k]}
        give = {k: min(room[k], int(s)) for k, s in shares.items()}
        if not any(give.values()):      # hand out the remainder by largest fraction
            for k in sorted(shares, key=lambda k: (-(shares[k] % 1), k))[:left]:
                give[k] = 1
        for k, g in give.items():
            g = min(g, left)
            alloc[k] += g
            room[k] -= g
            left -= g
    return alloc


def systematic(pool: list, n: int, rng: random.Random) -> list:
    """Systematic random sample: a random start, then every len/n-th item."""
    if n >= len(pool):
        return list(pool)
    step = len(pool) / n
    start = rng.random() * step
    return [pool[int(start + k * step)] for k in range(n)]


# --- data ------------------------------------------------------------------------

def companies(conn, labels: list[str]) -> list[dict]:
    ws.refresh_names(conn)
    out = []
    for w in ws.list_workspaces(conn, kinds=("real",)):
        label = client_label(w["slug"])
        if label in labels and not w.get("local_only"):
            uploads = repository.list_uploads(conn, source_client=w["slug"])
            if uploads:
                out.append({"slug": w["slug"], "label": label, "upload_id": max(u["id"] for u in uploads)})
    missing = set(labels) - {c["label"] for c in out}
    if missing:
        raise SystemExit(f"no hosted real workspace with an audit for: {', '.join(sorted(missing))}")
    return sorted(out, key=lambda c: labels.index(c["label"]))


def load_flags(conn, company: dict, data_root: Path) -> dict:
    """Re-run Layer 1 on the audit's files and collect every flagged row."""
    stored = {k for (k,) in conn.execute(
        "SELECT r.rule_key FROM issues i JOIN rules r ON r.id = i.rule_ref "
        "WHERE i.upload_id = %s AND i.source = 'rule'", (company["upload_id"],)).fetchall()}
    res = run_rules(data_root / company["slug"], source_client=company["slug"], collect_rows=True)
    fresh = {i["rule_key"]: i for i in res["issues"] if i.get("rule_key") in stored}
    if set(fresh) != stored:
        raise SystemExit(f"{company['label']}: the data folder no longer matches audit #{company['upload_id']}")
    tables = res["tables"]
    flagged = {t: set() for t in tables}
    wide = {t: [] for t in tables}
    for i in fresh.values():
        rows = i.get("_rows") or []
        if len(rows) > COLUMN_WIDE * len(tables[i["table_name"]]):
            wide[i["table_name"]].append(i["description"])
        else:
            flagged[i["table_name"]].update(rows)
    for table, ref in conn.execute(
            "SELECT table_name, row_reference FROM issues WHERE upload_id = %s AND source = 'llm'",
            (company["upload_id"],)).fetchall():
        if table in flagged:
            flagged[table].update((ref or {}).get("sample_rows") or [])
    return {"tables": tables, "flagged": flagged, "wide": wide}


def _cell(v) -> str:
    return "(empty)" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def sample_issues(conn, comps: list[dict], seed: int) -> list[dict]:
    rng = random.Random(seed)
    pools = {"rule": [], "llm": []}
    with conn.cursor(row_factory=dict_row) as cur:
        for c in comps:
            for i in cur.execute(
                    "SELECT i.*, r.what_text, r.why_text FROM issues i LEFT JOIN rules r ON r.id = i.rule_ref "
                    "WHERE i.upload_id = %s ORDER BY i.id", (c["upload_id"],)).fetchall():
                pools[i["source"]].append({**i, "company": c})
    n_ai = min(N_PER_SOURCE, len(pools["llm"]))
    n_rule = min(N_TRIAGE - n_ai, len(pools["rule"]))
    picked = rng.sample(pools["rule"], n_rule) + rng.sample(pools["llm"], n_ai)
    rng.shuffle(picked)
    print(f"triage sample: {n_rule} rule + {n_ai} AI issues")
    return picked


def sample_rows(flags: dict[str, dict], comps: list[dict], seed: int) -> list[dict]:
    """200 unflagged rows (stratified, systematic) and the fixed 50-row reviewer subset."""
    rng = random.Random(seed)
    pools = {}
    for c in comps:
        f = flags[c["slug"]]
        for t, df in f["tables"].items():
            pools[(c["slug"], t)] = [r for r in range(len(df)) if r not in f["flagged"][t]]
    alloc = allocate({k: len(v) for k, v in pools.items()}, N_RECALL, MIN_PER_TABLE)
    chosen = {k: systematic(pools[k], n, rng) for k, n in alloc.items() if n}
    sub_alloc = allocate({k: len(v) for k, v in chosen.items()}, N_SUBSET, 1)
    subset = {k: set(systematic(v, sub_alloc.get(k, 0), rng)) for k, v in chosen.items()}
    rows, no = [], 0
    for c in comps:
        for (slug, t), picked in chosen.items():
            if slug != c["slug"]:
                continue
            for r in picked:
                no += 1
                rows.append({"no": f"R{no:03d}", "slug": slug, "label": c["label"], "table": t,
                             "row": r, "subset": r in subset[(slug, t)], "upload_id": c["upload_id"]})
    return rows


# --- sheets ------------------------------------------------------------------------

def _header(ws_, cols: list[tuple[str, int]], row: int = 1) -> None:
    for j, (name, width) in enumerate(cols, start=1):
        cell = ws_.cell(row=row, column=j, value=name)
        cell.font, cell.fill, cell.alignment = BOLD, HEAD_FILL, WRAP
        ws_.column_dimensions[get_column_letter(j)].width = width
    ws_.freeze_panes = ws_.cell(row=row + 1, column=1)


def _dropdown(ws_, col: int, first: int, last: int, options: list[str]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws_.add_data_validation(dv)
    dv.add(f"{get_column_letter(col)}{first}:{get_column_letter(col)}{last}")


def _declaration(ws_, last_row: int) -> None:
    r = last_row + 3                                    # two empty rows above
    ws_.cell(row=r, column=1, value="REVIEWER DECLARATION").font = BOLD
    lines = [
        "I confirm that I reviewed these items myself, without seeing the first reviewer's decisions,",
        "and did not discuss individual items with anyone before finishing. My decisions reflect my own judgement.",
        "I will keep this data confidential, and I agree that my decisions may be used anonymously in the dissertation.",
    ]
    for k, line in enumerate(lines, start=1):
        ws_.cell(row=r + k, column=1, value=line)
    fields = ["Name:", "Role / job title:", "Years of ERP / business-data experience:", "Date:",
              "Signature (type your name):"]
    for k, label in enumerate(fields):
        row = r + len(lines) + 2 + k
        ws_.cell(row=row, column=1, value=label).font = BOLD
        ws_.cell(row=row, column=2).fill = FILL


def _start_tab(wb: Workbook, title: str, lines: list[str]) -> None:
    s = wb.active
    s.title = "Start here"
    s.column_dimensions["A"].width = 120
    s.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    for k, line in enumerate(lines, start=3):
        c = s.cell(row=k, column=1, value=line)
        c.alignment = WRAP


DIMENSION_LINES = [
    "The six kinds of problem:",
    "  Completeness - a value that should be there is missing.",
    "  Validity - a value is not allowed or impossible (e.g. a negative quantity, a wrong code).",
    "  Consistency - values in the same row contradict each other (e.g. total is not price x quantity).",
    "  Uniqueness - something that should appear once appears more than once.",
    "  Timeliness - a date is in the future or far out of date.",
    "  Cross-table - a value does not match the record it refers to in another table.",
]


def write_triage(path: Path, reviewer: str, items: list[dict], aliases: dict, tables: dict, masker: Masker) -> None:
    wb = Workbook()
    _start_tab(wb, f"Triage agreement - reviewer {reviewer}", [
        "Thank you for helping. This sheet lists 30 possible data quality problems that an automated check "
        "found in two anonymised companies' business data.",
        "For each issue, decide whether it is a real data quality problem in this data. Accept = real problem. "
        "Reject = not a problem (normal business data or a wrong finding). Unsure = cannot tell from the "
        "information given.",
        "Fill in only the shaded cells on the 'Triage' tab: Decision (drop-down) and, if you like, a short Reason.",
        "Values that could identify people are replaced by codes like MASK_1a2b3c4d. The same code always "
        "means the same original value.",
        *DIMENSION_LINES,
        f"Please finish by {DEADLINE}, then fill in the declaration under the last row.",
        "Please do not discuss individual items with anyone until you have finished.",
    ])
    s = wb.create_sheet("Triage")
    _header(s, [("No.", 6), ("Company", 12), ("Table", 18), ("Column", 22), ("What was found", 48),
                ("Why it was flagged (evidence)", 44), ("Example values (masked)", 44), ("Suggested fix", 40),
                ("Decision", 12), ("Reason (optional)", 36)])
    for n, i in enumerate(items, start=1):
        c = i["company"]
        names = aliases[c["slug"]].get(i["table_name"], {})
        cols = [x.strip() for x in (i["column_name"] or "").split(",") if x.strip()]
        column = ", ".join(f"{names[x]} ({x})" if x in names else x for x in cols) or "(whole row)"
        if i["source"] == "rule":
            evidence = " ".join(x for x in (i.get("what_text"), i.get("why_text")) if x)
        else:
            evidence = "Based on the column's values compared with the rest of the table."
        count = (i.get("row_reference") or {}).get("count") or i.get("affected_row_count")
        if count:
            evidence += f" Affected rows: {count:,}."
        values = _examples(i, cols, tables[c["slug"]], names, masker)
        row = [n, c["label"], i["table_name"], column, i["description"], evidence, values,
               i["recommendation"] or ""]
        for j, v in enumerate(row, start=1):
            s.cell(row=n + 1, column=j, value=v).alignment = WRAP
        for j in (9, 10):
            s.cell(row=n + 1, column=j).fill = FILL
    _dropdown(s, 9, 2, len(items) + 1, DECISIONS)
    _declaration(s, len(items) + 1)
    wb.save(path)


def _examples(issue: dict, cols: list[str], tables: dict, names: dict, masker: Masker) -> str:
    ref = issue.get("row_reference") or {}
    df = tables.get(issue["table_name"])
    rows = (ref.get("sample_rows") or [])[:N_EXAMPLES]
    if df is None or not rows:
        return ""
    show = cols or list(df.columns[:4])
    out = []
    for r in rows:
        if not 0 <= r < len(df):
            continue
        parts = []
        for col in show:
            if col in df.columns:
                v = df.at[r, col]
                v = None if pd.isna(v) else str(v)
                shown = _cell(masker.value(issue["upload_id"], issue["table_name"], col, v))
                parts.append(f"{names.get(col, col)} = {shown}")
        out.append(f"Row {r + 1}: " + "; ".join(parts))
    return "\n".join(out)


def write_recall(path: Path, who: str, rows: list[dict], flags: dict, aliases: dict, docs: list[dict],
                 masker: Masker, mark_subset: bool) -> None:
    n_rows = len(rows)
    wb = Workbook()
    _start_tab(wb, f"Recall check - {who}", [
        f"This workbook has {n_rows} rows from two anonymised companies' business data, one tab per table. "
        "The automated checks did not flag any of these rows.",
        "For each row, check whether anything looks wrong (missing, invalid, inconsistent, duplicated, out of "
        "date, or not matching another table). If yes, choose Y and the dimension.",
        "At the top of each tab, 'Already reported' lists problems that were found for the whole column - "
        "please do not mark those again.",
        "The 'Columns' tab explains what each column means. Values that could identify people are replaced by "
        "codes like MASK_1a2b3c4d; the same code always means the same original value.",
        *DIMENSION_LINES,
        "Fill in only the shaded cells: Issue found? (Y/N), Dimension (if Y) and an optional Note.",
        f"Please finish by {DEADLINE}, then fill in the declaration under the last row of each tab.",
        "Please do not discuss individual items with anyone until you have finished.",
    ])
    by_tab: dict[tuple, list[dict]] = {}
    for r in rows:
        by_tab.setdefault((r["label"], r["slug"], r["table"]), []).append(r)
    for (label, slug, table), items in by_tab.items():
        df = flags[slug]["tables"][table]
        names = aliases[slug].get(table, {})
        s = wb.create_sheet(f"{label[-1]} {table}"[:31])
        s.cell(row=1, column=1, value=f"{label} - {table}: {len(items)} rows to check").font = BOLD
        wide = flags[slug]["wide"][table]
        note = "Already reported for the whole column (do not mark again): " + (
            " | ".join(wide) if wide else "nothing")
        s.cell(row=2, column=1, value=note).alignment = WRAP
        s.row_dimensions[2].height = 15 * min(8, 1 + len(note) // 150)
        cols = [("No.", 7)] + [(f"{names[c]} ({c})" if c in names else c, 16) for c in df.columns]
        extra = [("Issue found?", 10), ("Dimension", 14), ("Note", 30)]
        if mark_subset:
            extra.append(("In reviewer subset", 10))
        _header(s, cols + extra, row=4)
        s.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(len(cols) + len(extra), 12))
        first = 5
        for k, r in enumerate(items):
            row = first + k
            s.cell(row=row, column=1, value=r["no"])
            for j, col in enumerate(df.columns, start=2):
                v = df.at[r["row"], col]
                v = None if pd.isna(v) else str(v)
                s.cell(row=row, column=j, value=_cell(masker.value(r["upload_id"], table, col, v))).alignment = WRAP
            base = len(cols)
            for j in (1, 2, 3):
                s.cell(row=row, column=base + j).fill = FILL
            if mark_subset:
                s.cell(row=row, column=base + 4, value="Y" if r["subset"] else "")
        last = first + len(items) - 1
        _dropdown(s, len(cols) + 1, first, last, ["Y", "N"])
        _dropdown(s, len(cols) + 2, first, last, DIMENSIONS)
        _declaration(s, last)
    c = wb.create_sheet("Columns")
    _header(c, [("Company", 12), ("Table", 20), ("Column", 12), ("Name", 26), ("What it means", 70),
                ("Status", 18)])
    for k, d in enumerate(docs, start=2):
        for j, v in enumerate([d["label"], d["table"], d["column"], d["name"], d["description"], d["status"]],
                              start=1):
            c.cell(row=k, column=j, value=v).alignment = WRAP
    wb.save(path)


def column_docs(conn, comps: list[dict], flags: dict) -> list[dict]:
    out = []
    for c in comps:
        docs = {(t, col): (name, desc, conf) for t, col, name, desc, conf in conn.execute(
            "SELECT table_name, column_name, display_name, description, confidence FROM schema_doc_columns "
            "WHERE source_client = %s", (c["slug"],)).fetchall()}
        for t, df in flags[c["slug"]]["tables"].items():
            for col in df.columns:
                name, desc, conf = docs.get((t, col), (None, None, None))
                out.append({"label": c["label"], "table": t, "column": col, "name": name or "",
                            "description": desc or "No description yet.",
                            "status": "verified" if conf == "VERIFIED" else ("AI draft (unverified)" if desc else "")})
    return out


# --- keys ---------------------------------------------------------------------------

def _write_key(path: Path, header: list[str], rows: list[list]) -> None:
    """Write an answer key; refuse if an existing key differs (data changed between runs)."""
    lines = [header] + [[str(x) for x in r] for r in rows]
    if path.exists():
        with path.open(newline="", encoding="utf-8") as f:
            if list(csv.reader(f)) != lines:
                raise SystemExit(f"{path.name} already exists with different content -- the sample changed. "
                                 "Move the old keys away only if no reviewer has received a sheet yet.")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(lines)


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--reviewer", required=True, help="e.g. R2")
    ap.add_argument("--seed", type=int, default=20261014)
    ap.add_argument("--companies", default="Company A,Company B")
    ap.add_argument("--data-root", default=str(ROOT / "data"))
    ap.add_argument("--out", default=str(OUT_DEFAULT))
    ap.add_argument("--keys", default=str(KEYS_DIR))
    args = ap.parse_args(argv)
    out, keys = Path(args.out), Path(args.keys)
    out.mkdir(parents=True, exist_ok=True)

    conn = repository.connect()
    comps = companies(conn, [x.strip() for x in args.companies.split(",")])
    masker = Masker(conn)
    aliases = {c["slug"]: display_aliases(conn, c["slug"]) for c in comps}
    flags = {}
    for c in comps:
        print(f"{c['label']}: audit #{c['upload_id']} -- re-running the rules to find every flagged row...")
        flags[c["slug"]] = load_flags(conn, c, Path(args.data_root))

    items = sample_issues(conn, comps, args.seed)
    rows = sample_rows(flags, comps, args.seed)
    _write_key(keys / "triage_key.csv", ["no", "issue_id", "company", "source"],
               [[n, i["id"], i["company"]["label"], i["source"]] for n, i in enumerate(items, start=1)])
    _write_key(keys / "recall_key.csv", ["no", "company", "table", "row", "in_subset"],
               [[r["no"], r["label"], r["table"], r["row"], "Y" if r["subset"] else "N"] for r in rows])

    tables = {c["slug"]: flags[c["slug"]]["tables"] for c in comps}
    docs = column_docs(conn, comps, flags)
    write_triage(out / f"Triage_Agreement_{args.reviewer}.xlsx", args.reviewer, items, aliases, tables, masker)
    write_recall(out / f"Recall_Check_{args.reviewer}.xlsx", f"reviewer {args.reviewer}",
                 [r for r in rows if r["subset"]], flags, aliases, docs, masker, mark_subset=False)
    mhg = out / "Recall_Check_MHG.xlsx"
    if not mhg.exists():
        write_recall(mhg, "first reviewer (MHG)", rows, flags, aliases, docs, masker, mark_subset=True)
    per_table = {}
    for r in rows:
        per_table[(r["label"], r["table"])] = per_table.get((r["label"], r["table"]), 0) + 1
    print("recall rows per table:", ", ".join(f"{l[-1]}/{t}: {n}" for (l, t), n in per_table.items()))
    print(f"subset for reviewers: {sum(r['subset'] for r in rows)} rows")
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
