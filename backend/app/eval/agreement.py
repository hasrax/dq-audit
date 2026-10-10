"""Score the second-reviewer sheets (Section 5.6): triage agreement + recall check.

Usage (from backend/):
    python -m app.eval.agreement --sheets Triage_Agreement_R2.xlsx [Triage_Agreement_R3.xlsx]
        --recall Recall_Check_MHG.xlsx Recall_Check_R2.xlsx [Recall_Check_R3.xlsx]
        [--out ../evaluation/results]

File names are looked up in evaluation/reviewer_packs/ when not found as
given (download the completed Google Sheets as .xlsx). MHG's triage
decisions come from the database (accepted / rejected / considered =
Accept / Reject / Unsure; issues still "to review" count as missing).

Triage: percent agreement and Cohen's kappa for MHG vs each reviewer (and
R2 vs R3), Fleiss' kappa when there are three raters; overall and per
detector (rule / AI); Unsure kept as its own category and, separately,
excluded. Plus a confusion table and a disagreement list.
Recall: MHG's miss rate on the 200 rows and each reviewer's on the
50-row subset (Wilson 95% CI), and MHG-vs-reviewer agreement on the
overlapping rows. Reviewer names never appear in outputs; only role and
years of experience from the declaration. Every input file's SHA-256 is
printed and written to agreement_inputs.json for the snapshot manifest.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.db import repository

ROOT = Path(__file__).resolve().parents[3]
PACKS = ROOT / "evaluation" / "reviewer_packs"
KEYS = ROOT / "evaluation" / "reviewer_keys"
OUT_DEFAULT = ROOT / "evaluation" / "results"
STATUS_TO_DECISION = {"accepted": "Accept", "rejected": "Reject", "considered": "Unsure"}
DECISIONS = ["Accept", "Reject", "Unsure"]


# --- statistics (pure, tested) ---------------------------------------------------

def cohen_kappa(a: list, b: list) -> float | None:
    """Cohen's kappa for two raters over the same items (None if undefined)."""
    n = len(a)
    if n == 0:
        return None
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return None if pe == 1 else (po - pe) / (1 - pe)


def fleiss_kappa(ratings: list[list]) -> float | None:
    """Fleiss' kappa; `ratings` is one list of category labels per item (same rater count)."""
    if not ratings:
        return None
    m = len(ratings[0])
    cats = sorted({c for item in ratings for c in item})
    n = len(ratings)
    p_i = [(sum(cnt * cnt for cnt in Counter(item).values()) - m) / (m * (m - 1)) for item in ratings]
    p_bar = sum(p_i) / n
    p_j = [sum(item.count(c) for item in ratings) / (n * m) for c in cats]
    pe = sum(p * p for p in p_j)
    return None if pe == 1 else (p_bar - pe) / (1 - pe)


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    if n == 0:
        return (0.0, 0.0)
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return (max(0.0, centre - half), min(1.0, centre + half))


def pairwise(a: dict, b: dict, items: list, exclude_unsure: bool = False) -> dict:
    keep = [i for i in items if a.get(i) and b.get(i)
            and not (exclude_unsure and "Unsure" in (a[i], b[i]))]
    x, y = [a[i] for i in keep], [b[i] for i in keep]
    k = cohen_kappa(x, y)
    return {"n": len(keep), "agreement_pct": round(100 * sum(p == q for p, q in zip(x, y)) / len(keep), 1)
            if keep else None, "kappa": round(k, 3) if k is not None else None}


# --- reading sheets ------------------------------------------------------------------

def _resolve(name: str) -> Path:
    p = Path(name)
    return p if p.exists() else PACKS / name


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _declaration(wb) -> dict:
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                label = str(c.value or "")
                nxt = ws.cell(row=c.row, column=c.column + 1).value
                if label.startswith("Role / job title") and nxt and "role" not in out:
                    out["role"] = str(nxt).strip()
                if label.startswith("Years of ERP") and nxt and "years" not in out:
                    out["years"] = str(nxt).strip()
    return out


def read_triage(path: Path) -> tuple[dict[int, str], dict[int, str], dict]:
    wb = load_workbook(path, read_only=False)
    ws = wb["Triage"]
    decisions, reasons = {}, {}
    for row in ws.iter_rows(min_row=2):
        no = row[0].value
        if isinstance(no, (int, float)) or (isinstance(no, str) and no.isdigit()):
            d = str(row[8].value or "").strip().capitalize()
            if d in DECISIONS:
                decisions[int(no)] = d
            if row[9].value:
                reasons[int(no)] = str(row[9].value)
    return decisions, reasons, _declaration(wb)


def read_recall(path: Path) -> tuple[dict[str, dict], dict]:
    wb = load_workbook(path, read_only=False)
    out = {}
    for ws in wb.worksheets:
        header = [str(c.value or "") for c in ws[4]] if ws.max_row >= 4 else []
        if "Issue found?" not in header:
            continue
        f, d, n = (header.index(h) for h in ("Issue found?", "Dimension", "Note"))
        for row in ws.iter_rows(min_row=5):
            no = str(row[0].value or "")
            if re.fullmatch(r"R\d{3}", no):
                found = str(row[f].value or "").strip().upper()[:1]
                out[no] = {"found": found if found in ("Y", "N") else None,
                           "dimension": row[d].value, "note": row[n].value}
    return out, _declaration(wb)


def _reviewer(path: Path) -> str:
    m = re.search(r"_(R\d+|MHG)\.xlsx$", path.name, re.I)
    return m.group(1).upper() if m else path.stem


def mhg_decisions(conn, key: list[dict]) -> dict[int, str]:
    ids = [int(k["issue_id"]) for k in key]
    status = dict(conn.execute("SELECT id, status FROM issues WHERE id = ANY(%s)", (ids,)).fetchall())
    return {int(k["no"]): STATUS_TO_DECISION.get(status.get(int(k["issue_id"]))) for k in key}


def _read_key(name: str) -> list[dict]:
    with (KEYS / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --- report ------------------------------------------------------------------------

def _table(rows: list[dict]) -> str:
    if not rows:
        return "_none_\n"
    cols = list(rows[0])
    lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    lines += ["| " + " | ".join("" if r[c] is None else str(r[c]) for c in cols) + " |" for r in rows]
    return "\n".join(lines) + "\n"


def score_triage(raters: dict[str, dict[int, str]], key: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    source = {int(k["no"]): ("AI" if k["source"] == "llm" else "Rule") for k in key}
    items = sorted(source)
    names = list(raters)
    rows = []
    pairs = [("MHG", r) for r in names if r != "MHG"] + [
        (a, b) for i, a in enumerate(names) for b in names[i + 1:] if "MHG" not in (a, b)]
    for a, b in pairs:
        for scope in ("All", "Rule", "AI"):
            its = [i for i in items if scope == "All" or source[i] == scope]
            for excl in (False, True):
                rows.append({"raters": f"{a} vs {b}", "detector": scope,
                             "unsure": "excluded" if excl else "own category",
                             **pairwise(raters[a], raters[b], its, excl)})
    if len(names) >= 3:
        full = [i for i in items if all(raters[r].get(i) for r in names)]
        k = fleiss_kappa([[raters[r][i] for r in names] for i in full])
        rows.append({"raters": " + ".join(names) + " (Fleiss)", "detector": "All", "unsure": "own category",
                     "n": len(full), "agreement_pct": None, "kappa": round(k, 3) if k is not None else None})
    confusion, disagreements = [], []
    for r in names:
        if r == "MHG":
            continue
        for d in DECISIONS:
            confusion.append({"reviewer": r, "MHG": d, **{f"{r} {e}": sum(
                1 for i in items if raters["MHG"].get(i) == d and raters[r].get(i) == e) for e in DECISIONS}})
        for i in items:
            if raters["MHG"].get(i) and raters[r].get(i) and raters["MHG"][i] != raters[r][i]:
                disagreements.append({"no": i, "issue_id": next(k["issue_id"] for k in key if int(k["no"]) == i),
                                      "detector": source[i], "MHG": raters["MHG"][i], r: raters[r][i]})
    return rows, confusion, disagreements


def score_recall(sheets: dict[str, dict], key: list[dict]) -> tuple[list[dict], list[dict]]:
    subset = {k["no"] for k in key if k["in_subset"] == "Y"}
    rows, agree = [], []
    for who, answers in sheets.items():
        scope = [k["no"] for k in key] if who == "MHG" else sorted(subset)
        rated = [n for n in scope if answers.get(n, {}).get("found")]
        misses = sum(answers[n]["found"] == "Y" for n in rated)
        lo, hi = wilson(misses, len(rated))
        dims = Counter(answers[n]["dimension"] for n in rated if answers[n]["found"] == "Y")
        rows.append({"rater": who, "rows_in_sheet": len(scope), "rows_rated": len(rated), "issues_found": misses,
                     "miss_rate_pct": round(100 * misses / len(rated), 1) if rated else None,
                     "ci95_low_pct": round(100 * lo, 1), "ci95_high_pct": round(100 * hi, 1),
                     "by_dimension": "; ".join(f"{d}: {c}" for d, c in dims.most_common()) or ""})
    if "MHG" in sheets:
        for who, answers in sheets.items():
            if who == "MHG":
                continue
            a = {n: sheets["MHG"].get(n, {}).get("found") for n in subset}
            b = {n: answers.get(n, {}).get("found") for n in subset}
            agree.append({"raters": f"MHG vs {who}", **pairwise(a, b, sorted(subset))})
    return rows, agree


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sheets", nargs="+", required=True)
    ap.add_argument("--recall", nargs="*", default=[])
    ap.add_argument("--out", default=str(OUT_DEFAULT))
    args = ap.parse_args(argv)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    inputs, profiles = [], []

    triage_key = _read_key("triage_key.csv")
    conn = repository.connect()
    raters = {"MHG": mhg_decisions(conn, triage_key)}
    reasons = {}
    for name in args.sheets:
        path = _resolve(name)
        who = _reviewer(path)
        raters[who], reasons[who], decl = read_triage(path)
        inputs.append(path)
        profiles.append(f"{who}: {decl.get('role', 'role not given')}, {decl.get('years', '?')} years of "
                        "ERP / business-data experience")
    missing = {who: len([k for k in triage_key if not raters[who].get(int(k["no"]))]) for who in raters}
    rows, confusion, disagreements = score_triage(raters, triage_key)
    for d in disagreements:
        for who in raters:
            if who != "MHG" and who in d:
                d["reason"] = reasons[who].get(d["no"], "")

    recall_rows, recall_agree = [], []
    if args.recall:
        recall_key = _read_key("recall_key.csv")
        sheets = {}
        for name in args.recall:
            path = _resolve(name)
            sheets[_reviewer(path)], _ = read_recall(path)
            inputs.append(path)
        recall_rows, recall_agree = score_recall(sheets, recall_key)

    with (out / "agreement.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    md = ["# Triage agreement (T4)", "", "Reviewer profiles: " + "; ".join(profiles), "",
          "Items without a decision: " + ", ".join(f"{k} {v}" for k, v in missing.items()), "",
          _table(rows), "## Confusion (rows = MHG)", "", _table(confusion), "## Disagreements", "",
          _table(disagreements)]
    (out / "agreement.md").write_text("\n".join(md), encoding="utf-8")
    if recall_rows:
        with (out / "recall.csv").open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(recall_rows[0]))
            w.writeheader()
            w.writerows(recall_rows)
        (out / "recall.md").write_text("\n".join([
            "# Recall check (T5)", "", "Miss rate = rows a rater found a problem in, among rows the system did "
            "not flag (column-wide findings excluded, majority rule). Wilson 95% CI.", "", _table(recall_rows),
            "## Agreement on the reviewer subset", "", _table(recall_agree)]), encoding="utf-8")

    stamp = datetime.now().isoformat(timespec="seconds")
    manifest = [{"file": p.name, "sha256": sha256(p), "scored_at": stamp} for p in inputs]
    (out / "agreement_inputs.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    for m in manifest:
        print(f"{m['sha256']}  {m['file']}")
    print(_table([r for r in rows if r["detector"] == "All" and r["unsure"] == "own category"]))
    if recall_rows:
        print(_table(recall_rows))
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
