"""CSV and Excel export of audit results (Layer 3).

Pure builders (no I/O) so they are testable and reusable by the
dashboard's download buttons and any CLI. Exports carry the anonymised
organisation label (Company A/B/C), never the real client name, and
column names are aliased for readability where an alias map is given.
"""
from __future__ import annotations

import csv
import io
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.aliases import display_columns
from app.reports.wording import DATE_FORMAT, show

SOURCE_LABEL = {"rule": "Rule", "llm": "AI"}

ISSUE_FIELDS = [
    "organisation", "upload", "table", "column", "dimension", "source", "severity",
    "trend_status", "affected_rows", "explanation", "recommendation", "fixable",
    "status", "reviewer", "reviewed_at", "user_note",
]
SCORE_FIELDS = [
    "overall_score", "completeness_score", "validity_score", "consistency_score",
    "uniqueness_score", "timeliness_score", "cross_table_score",
]
SEVERITIES = ["high", "medium", "low"]
WIDE_COLUMNS = {"explanation": 70, "recommendation": 50, "user_note": 30}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def issue_rows(issues: list[dict], aliases: dict | None = None, organisation: str = "") -> list[dict]:
    """One flat, export-ready dict per issue in ISSUE_FIELDS order."""
    aliases = aliases or {}
    rows = []
    for issue in issues:
        ref = issue.get("row_reference") or {}
        affected = issue.get("affected_row_count")
        reviewed_at = issue.get("reviewed_at")
        rows.append({
            "organisation": organisation,
            "upload": issue.get("upload_id"),
            "table": issue["table_name"],
            "column": display_columns(aliases, issue["table_name"], issue.get("column_name")) or "",
            "dimension": issue["dimension"],
            "source": SOURCE_LABEL.get(issue["source"], issue["source"]),
            "severity": issue["severity"],
            "trend_status": issue.get("trend_status") or "",
            "affected_rows": affected if affected is not None else ref.get("count"),
            "explanation": show(issue.get("description"), issue["table_name"], aliases),
            "recommendation": show(issue.get("recommendation"), issue["table_name"], aliases),
            "fixable": issue.get("fixable"),
            "status": issue.get("status"),
            "reviewer": issue.get("reviewed_by") or "",
            "reviewed_at": (reviewed_at.strftime(f"{DATE_FORMAT} %H:%M") if isinstance(reviewed_at, datetime)
                            else (reviewed_at or "")),
            "user_note": issue.get("user_note") or "",
        })
    return rows


def issues_to_csv(issues: list[dict], aliases: dict | None = None, organisation: str = "") -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=ISSUE_FIELDS)
    writer.writeheader()
    writer.writerows(issue_rows(issues, aliases, organisation))
    return buf.getvalue()


def issues_to_xlsx(
    issues: list[dict], aliases: dict | None = None, organisation: str = "",
    title: str = "", filters: dict | None = None,
) -> bytes:
    """Formatted workbook: an Issues sheet (styled header, frozen header
    row, auto-filter, readable widths) and a Summary sheet with counts by
    dimension and severity."""
    rows = issue_rows(issues, aliases, organisation)
    wb = Workbook()

    ws = wb.active
    ws.title = "Issues"
    ws.append([f.replace("_", " ").title() for f in ISSUE_FIELDS])
    for r in rows:
        ws.append([r[f] for f in ISSUE_FIELDS])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ISSUE_FIELDS))}{max(len(rows) + 1, 1)}"
    wrap = Alignment(wrap_text=True, vertical="top")
    for idx, field in enumerate(ISSUE_FIELDS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = WIDE_COLUMNS.get(field, max(12, len(field) + 4))
        if field in WIDE_COLUMNS:
            for cell in ws[letter][1:]:
                cell.alignment = wrap

    summary = wb.create_sheet("Summary")
    summary.append([title or "Data quality audit — issue export"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Organisation: {organisation}" if organisation else ""])
    summary.append([f"Exported {datetime.now():{DATE_FORMAT} %H:%M} · {len(rows)} issues"])
    active = {k: v for k, v in (filters or {}).items() if v not in (None, "", "(all)")}
    summary.append(["Filters: " + (", ".join(f"{k} = {v}" for k, v in active.items()) if active else "none")])
    summary.append([])

    header_row = summary.max_row + 1
    summary.append(["Dimension", *[s.title() for s in SEVERITIES], "Total"])
    counts = Counter((r["dimension"], r["severity"]) for r in rows)
    for dim in sorted({r["dimension"] for r in rows}):
        line = [counts[(dim, s)] for s in SEVERITIES]
        summary.append([dim, *line, sum(line)])
    totals = [sum(counts[(d, s)] for d in {r["dimension"] for r in rows}) for s in SEVERITIES]
    summary.append(["Total", *totals, sum(totals)])
    for cell in summary[header_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)

    summary.append([])
    summary.append(["By source", "Issues"])
    for cell in summary[summary.max_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for source, n in sorted(Counter(r["source"] for r in rows).items()):
        summary.append([source, n])
    summary.column_dimensions["A"].width = 28
    for letter in "BCDE":
        summary.column_dimensions[letter].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def score_to_csv(score: dict) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["metric", "value"])
    for field in SCORE_FIELDS:
        writer.writerow([field, score.get(field)])
    return buf.getvalue()
