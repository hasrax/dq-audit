import io

import pytest
from docx import Document

from app.db import repository
from app.reports.audit_report import build_markdown
from app.reports.render import parse, to_docx, to_pdf

SAMPLE = """# Title

Intro with **bold** text — and a dash.

## Section

- first point
- second point

| A | B |
|---|---|
| 1 | 2 |
"""


def test_parse_recognises_blocks():
    kinds = [b[0] for b in parse(SAMPLE)]
    assert kinds == ["h", "p", "h", "li", "li", "table"]
    assert parse(SAMPLE)[-1][1] == [["A", "B"], ["1", "2"]]


def test_docx_and_pdf_render():
    doc = Document(io.BytesIO(to_docx(SAMPLE)))
    assert doc.paragraphs[0].text == "Title" and len(doc.tables) == 1
    pdf = to_pdf(SAMPLE)
    assert pdf.startswith(b"%PDF") and len(pdf) > 500


def test_standard_report_builds_for_a_real_upload():
    try:
        conn = repository.connect()
    except Exception as exc:  # noqa: BLE001
        pytest.skip(f"PostgreSQL not available: {exc}")
    uploads = repository.list_uploads(conn)
    if not uploads:
        pytest.skip("no uploads")
    md = build_markdown(conn, uploads[0])
    assert md.startswith("# Data quality audit report")
    for section in ("## 1. Executive summary", "## 2. Scope", "## 3. Method", "## 4. Results",
                    "## 5. Key findings", "## 6. Recommendations", "## Appendix A. Definitions",
                    "### 4.3 What to do first"):
        assert section in md
    assert to_pdf(md).startswith(b"%PDF")
