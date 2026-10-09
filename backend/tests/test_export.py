import csv
import io

from app.export import issues_to_csv, score_to_csv


def _issue(**kw):
    base = {
        "table_name": "GRNHED", "column_name": "col_5", "dimension": "completeness",
        "source": "rule", "error_type": "type1", "severity": "high", "status": "flagged",
        "fixable": "no", "row_reference": {"count": 12, "sample_rows": [0, 1]},
        "description": "d", "recommendation": "r", "rule_id": "completeness_null_pct",
    }
    base.update(kw)
    return base


def test_issues_csv_has_header_and_rows():
    text = issues_to_csv([_issue(), _issue(severity="low")])
    rows = list(csv.DictReader(io.StringIO(text)))
    assert len(rows) == 2
    assert rows[0]["dimension"] == "completeness"
    assert rows[0]["affected_rows"] == "12"


def test_issues_csv_applies_aliases():
    aliases = {"GRNHED": {"col_5": "Invoice No"}}
    text = issues_to_csv([_issue()], aliases)
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows[0]["column"] == "Invoice No"


def test_issues_csv_handles_null_column_and_row_reference():
    text = issues_to_csv([_issue(column_name=None, row_reference=None)])
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows[0]["column"] == ""
    assert rows[0]["affected_rows"] == ""


def test_score_csv_lists_all_dimensions():
    score = {
        "overall_score": 30.4, "completeness_score": 10.4, "validity_score": 34.5,
        "consistency_score": 28.6, "uniqueness_score": 52.6, "timeliness_score": 62.5,
        "cross_table_score": 25.6,
    }
    rows = list(csv.reader(io.StringIO(score_to_csv(score))))
    assert rows[0] == ["metric", "value"]
    assert ["overall_score", "30.4"] in rows
    assert len(rows) == 8  # header + 7 metrics


def test_xlsx_has_frozen_filtered_issue_sheet_and_summary():
    from openpyxl import load_workbook

    from app.export import issues_to_xlsx

    data = issues_to_xlsx([_issue(), _issue(severity="low", dimension="validity")],
                          organisation="Company A", filters={"status": "flagged", "source": "(all)"})
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Issues"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert ws["A2"].value == "Company A" and ws.max_row == 3
    summary = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert ["Total", 1, 0, 1, 2] in summary
    assert any(r[0] == "Filters: status = flagged" for r in summary)


def test_csv_uses_organisation_label_and_trend():
    text = issues_to_csv([_issue(trend_status="worsening", affected_row_count=40)], organisation="Company B")
    row = next(csv.DictReader(io.StringIO(text)))
    assert row["organisation"] == "Company B" and row["trend_status"] == "worsening" and row["affected_rows"] == "40"
