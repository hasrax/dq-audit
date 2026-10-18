"""One Excel workbook with every evaluation result, timestamped (evidence trail).

Usage (from backend/):
    python -m app.eval.export_workbook [--results ../evaluation/results] [--freeze]

Also scores the T11 schema-check sheet (evaluation/T11_schema_check_CompanyA.xlsx) if present.
Writes evaluation/results/evaluation_results.xlsx and a timestamped copy in
evaluation/results/workbooks/ (copies are never overwritten). The README tab records:
generation time (UTC and Sri Lanka time), git commit and branch, the AI model,
and for every source file its modified time and SHA-256, so a reader can check
that the numbers in the dissertation match the files they came from.
"""
from __future__ import annotations

import argparse
import hashlib
import subprocess
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import load_settings

ROOT = Path(__file__).resolve().parents[3]
RESULTS = ROOT / "evaluation" / "results"
T11_SHEET = ROOT / "evaluation" / "T11_schema_check_CompanyA.xlsx"
SL = timezone(timedelta(hours=5, minutes=30))
# tab name -> source CSV (in the results folder), in reading order
TABS = [
    ("T3 Precision", "precision.csv"), ("T5 Injection", "injection_coverage.csv"),
    ("T6 Reproducibility", "reproducibility.csv"),
    ("T7 AI stability", "ai_stability.csv"), ("T7 Pairs", "ai_stability_pairs.csv"),
    ("T8 Rules vs AI", "complementarity.csv"), ("T10 Projection", "projection_error.csv"),
    ("T11 Schema check", "schema_check_t11.csv"), ("T11 Schema docs", "schema_docs.csv"),
    ("T13 Tasks all", "usability_tasks.csv"), ("T13 Tasks v1", "usability_tasks_v1.csv"),
    ("T13 Tasks v2", "usability_tasks_v2.csv"), ("T13 Participants", "usability.csv"),
    ("T13 SUS", "sus_scores.csv"), ("T13 Integrity", "integrity.csv"), ("T14 Tests", "automated_tests.csv"),
    ("T15 Privacy", "privacy_prompts.csv"),
]


# Summary tabs copied into the app's "Thesis snapshot" view (per-participant tabs stay out). The file is
# local only: spreadsheets are git-ignored, so the hosted site shows a note instead.
SNAPSHOT = ROOT / "config" / "thesis_snapshot.xlsx"
SNAPSHOT_TABS = ["T3 Precision", "T5 Injection", "T6 Reproducibility", "T7 AI stability", "T7 Pairs",
                 "T8 Rules vs AI", "T10 Projection", "T11 Schema check", "T11 Schema docs", "T13 Tasks v1",
                 "T13 Tasks v2", "T15 Privacy"]


def freeze(workbook: Path, dst: Path = SNAPSHOT) -> Path:
    """Copy the summary tabs (plus the README facts and a SUS summary per study
    iteration) into the file the Evaluation page shows as the thesis snapshot."""
    sheets = pd.read_excel(workbook, sheet_name=None)
    with pd.ExcelWriter(dst, engine="openpyxl") as xw:
        readme = sheets["README"]
        readme.iloc[:6, :2].to_excel(xw, sheet_name="README", index=False)
        for tab in SNAPSHOT_TABS:
            if tab in sheets:
                sheets[tab].to_excel(xw, sheet_name=tab, index=False)
        sus = sheets.get("T13 SUS")
        if sus is not None and "version" in sus.columns:
            items = [c for c in ("sus", "C1", "C2", "C3", "C4", "C5") if c in sus.columns]
            rows = [{"iteration": v, "n": len(g), **g[items].mean().round(1).to_dict()}
                    for v, g in sus.groupby("version")]
            rows.append({"iteration": "All", "n": len(sus), **sus[items].mean().round(1).to_dict()})
            pd.DataFrame(rows).to_excel(xw, sheet_name="T13 SUS summary", index=False)
    return dst


def score_t11(sheet: Path, out: Path) -> pd.DataFrame | None:
    """T11 from the researcher's judgements: AI-drafted names/descriptions vs the source schema."""
    if not sheet.exists():
        return None
    df = pd.read_excel(sheet, sheet_name="Schema check")
    df = df[df["Judgement"].notna()]
    if df.empty:
        return None
    table = (df.groupby(["AI confidence", "Judgement"]).size().unstack(fill_value=0)
             .reindex(columns=["Correct", "Partly correct", "Wrong"], fill_value=0))
    table["n"] = table.sum(axis=1)
    total = table.sum().to_frame().T
    total.index = ["All"]
    table = pd.concat([table, total])
    for c in ("Correct", "Partly correct", "Wrong"):
        table[f"{c} %"] = (100 * table[c] / table["n"]).round(1)
    table = table.reset_index().rename(columns={"index": "AI confidence"})
    table.to_csv(out / "schema_check_t11.csv", index=False)
    return table


def _git(*args: str) -> str:
    try:
        return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, timeout=10).stdout.strip()
    except Exception:  # noqa: BLE001
        return "unknown"


def _write_df(ws, df: pd.DataFrame) -> None:
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=str(col))
        c.font, c.fill = Font(bold=True), PatternFill("solid", start_color="E8EEF4")
        ws.column_dimensions[get_column_letter(j)].width = min(45, max(10, len(str(col)) + 2))
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=None if pd.isna(v) else v).alignment = Alignment(wrap_text=True,
                                                                                          vertical="top")
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results", default=str(RESULTS))
    ap.add_argument("--freeze", action="store_true",
                    help=f"also write the thesis-snapshot summary for the Evaluation page ({SNAPSHOT.name})")
    args = ap.parse_args(argv)
    res = Path(args.results)
    t11 = score_t11(T11_SHEET, res)
    now = datetime.now(timezone.utc)
    settings = load_settings()

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    sources, rows = [], [
        ("Workbook", "ClarityDQ evaluation results (Chapter 5)"),
        ("Generated (UTC)", now.strftime("%Y-%m-%d %H:%M:%S")),
        ("Generated (Sri Lanka time)", now.astimezone(SL).strftime("%Y-%m-%d %H:%M:%S")),
        ("Git branch / commit", f"{_git('rev-parse', '--abbrev-ref', 'HEAD')} / {_git('rev-parse', '--short', 'HEAD')}"),
        ("AI model (thesis numbers)", f"{settings.llm_model} via {settings.llm_base_url}"),
        ("Note", "Every tab is a copy of a CSV produced by the evaluation scripts (backend/app/eval). "
                 "Hashes below let a reader check the files have not changed since this workbook was made."),
    ]
    for name, csv_name in TABS:
        path = res / csv_name
        if not path.exists():
            sources.append((name, csv_name, "not yet available", "", ""))
            continue
        df = pd.read_csv(path)
        _write_df(wb.create_sheet(name[:31]), df)
        mtime = datetime.fromtimestamp(path.stat().st_mtime, timezone.utc)
        sources.append((name, csv_name, f"{len(df)} rows", mtime.strftime("%Y-%m-%d %H:%M:%S UTC"),
                        hashlib.sha256(path.read_bytes()).hexdigest()))
    for i, (k, v) in enumerate(rows, start=1):
        readme.cell(row=i, column=1, value=k).font = Font(bold=True)
        readme.cell(row=i, column=2, value=v)
    start = len(rows) + 2
    for j, h in enumerate(["Tab", "Source file", "Contents", "File modified", "SHA-256"], start=1):
        readme.cell(row=start, column=j, value=h).font = Font(bold=True)
    for i, src in enumerate(sources, start=start + 1):
        for j, v in enumerate(src, start=1):
            readme.cell(row=i, column=j, value=v)
    for col, w in zip("ABCDE", (28, 30, 20, 24, 70)):
        readme.column_dimensions[col].width = w

    main_path = res / "evaluation_results.xlsx"
    wb.save(main_path)
    hist = res / "workbooks"
    hist.mkdir(exist_ok=True)
    copy = hist / f"evaluation_results_{now.astimezone(SL):%Y-%m-%d_%H%M%S}.xlsx"
    if not copy.exists():
        wb.save(copy)
    if t11 is not None:
        print(t11.to_string(index=False))
    print(f"tabs: {sum(1 for s in sources if s[2] != 'not yet available')} of {len(TABS)}; missing: "
          + ", ".join(s[0] for s in sources if s[2] == "not yet available"))
    print(f"wrote {main_path} and {copy.name}")
    if args.freeze:
        print(f"froze the thesis snapshot into {freeze(main_path)}")


if __name__ == "__main__":
    main()
